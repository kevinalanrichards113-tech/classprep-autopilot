#!/usr/bin/env python3
"""ClassPrep Autopilot — Builders.

Deck Builder (first Builder, pulled forward to Day 1): asks Muse Spark
for a visual-first deck spec — slides carry a kicker, a big headline and
an on-screen visual cue; the full narration script lives in the speaker
notes — then renders a real .pptx with python-pptx in the app's dark,
glossy style. A deterministic fallback spec is built from the plan and
curator items if the model's spec cannot be parsed.
"""
import json
import os
import re

import meta_api

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "outputs")

DECK_SYSTEM = (
    "You are the Deck Builder stage of ClassPrep Autopilot. Build a "
    "visual-first slide deck for one class session from the plan and "
    "curated sources. Rules: slides SUPPLEMENT the spoken narration, "
    "they do not replace it — each slide gets a short kicker, one big "
    "headline, and a brief on-screen visual cue; the full teaching "
    "script for that slide goes in 'narration' (spoken style, aimed at "
    "about 60-90 seconds). Narration must stand on its own: NO overt "
    "back-references like 'as we saw earlier' or 'in the last slide' — "
    "continuity comes from content, not signposting. Follow the "
    "session arc in the plan. Return JSON only: {\"slides\": "
    "[{\"kicker\": str, \"title\": str, \"visual\": str, "
    "\"narration\": str}]} with 6-9 slides. No prose outside the JSON."
)

BG = (11, 11, 18)
ACCENT = (110, 168, 255)
WHITE = (240, 242, 248)
DIM = (158, 166, 184)


def _as_json(value):
    if isinstance(value, (dict, list)):
        return value
    if isinstance(value, str):
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            m = re.search(r"(\{.*\}|\[.*\])", value, re.S)
            if m:
                try:
                    return json.loads(m.group(1))
                except json.JSONDecodeError:
                    return None
    return None


def _items_list(items):
    """Curator items as a list of dicts. The model sometimes returns the
    list double-encoded (each item a JSON string) — normalize."""
    out = []
    for it in _as_json(items) or []:
        if isinstance(it, dict):
            out.append(it)
        elif isinstance(it, str):
            parsed = _as_json(it)
            if isinstance(parsed, dict):
                out.append(parsed)
    return out


def _fallback_spec(session, plan, items):
    plan = _as_json(plan) or {}
    items = _as_json(items) or []
    slides = []
    for obj in (plan.get("objectives") or [])[:4]:
        slides.append({
            "kicker": "TODAY'S GOALS",
            "title": obj,
            "visual": "Goal statement on a clean dark slide.",
            "narration": f"By the end of this session, you will be able "
                         f"to do this: {obj}",
        })
    for seg in (plan.get("session_arc") or [])[:5]:
        slides.append({
            "kicker": (seg.get("segment") or "SEGMENT").upper(),
            "title": seg.get("focus") or seg.get("segment") or "Segment",
            "visual": "Headline with a single supporting visual.",
            "narration": seg.get("focus") or "",
        })
    for it in items[:3]:
        slides.append({
            "kicker": (it.get("type") or "SOURCE").upper(),
            "title": it.get("title") or "Source",
            "visual": f"Source: {it.get('source_name', '')}",
            "narration": it.get("why_it_fits") or "",
        })
    return {"slides": slides[:9]}


def deck_spec(session, plan, items):
    raw = meta_api.chat(
        DECK_SYSTEM,
        "Session input:\n" + json.dumps(session, indent=2)
        + "\n\nPlan:\n" + (plan if isinstance(plan, str)
                           else json.dumps(plan, indent=2))
        + "\n\nCurated items:\n" + (items if isinstance(items, str)
                                    else json.dumps(items, indent=2)),
        max_output_tokens=6000)
    spec = _as_json(raw)
    if not spec or not spec.get("slides"):
        spec = _fallback_spec(session, plan, items)
    return spec


def _blank(prs):
    return prs.slide_layouts[6]


def _add_text(slide, left, top, width, height, text, size, color,
              bold=False):
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    from pptx.dml.color import RGBColor
    p.font.color.rgb = RGBColor(*color)
    return box


def _bg(slide, prs):
    from pptx.dml.color import RGBColor
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*BG)


def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text or ""


def render_pptx(session, spec, items, slug):
    from pptx import Presentation
    from pptx.util import Inches

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Title slide
    s = prs.slides.add_slide(_blank(prs))
    _bg(s, prs)
    _add_text(s, 0.9, 1.6, 11.5, 0.6, (session.get("course") or "").upper(),
              20, ACCENT, bold=True)
    _add_text(s, 0.9, 2.3, 11.5, 1.8, session.get("topic") or "Session",
              48, WHITE, bold=True)
    _add_text(s, 0.9, 4.3, 11.5, 0.6, session.get("date") or "", 22, DIM)
    _add_text(s, 0.9, 6.5, 11.5, 0.5,
              "Built by ClassPrep Autopilot — DRAFT, awaiting instructor "
              "review", 13, DIM)
    _notes(s, "Deck generated by the ClassPrep Autopilot Deck Builder. "
              "Slides are visual-first; the teaching script is in the "
              "speaker notes of each slide.")

    # Content slides from the spec
    for sl in spec.get("slides", []):
        s = prs.slides.add_slide(_blank(prs))
        _bg(s, prs)
        _add_text(s, 0.9, 0.7, 11.5, 0.5, (sl.get("kicker") or "").upper(),
                  16, ACCENT, bold=True)
        _add_text(s, 0.9, 1.3, 11.5, 2.2, sl.get("title") or "", 36,
                  WHITE, bold=True)
        if sl.get("visual"):
            _add_text(s, 0.9, 4.6, 11.5, 1.6,
                      "On screen: " + sl["visual"], 18, DIM)
        _notes(s, sl.get("narration") or "")

    # Sources slide (every curator item, so the Verifier can check them)
    items = _items_list(items)
    if items:
        s = prs.slides.add_slide(_blank(prs))
        _bg(s, prs)
        _add_text(s, 0.9, 0.7, 11.5, 0.5, "SOURCES", 16, ACCENT, bold=True)
        lines = [f"{it.get('title', '')} — {it.get('source_name', '')}\n"
                 f"{it.get('source_url', '')}" for it in items]
        _add_text(s, 0.9, 1.4, 11.5, 5.4, "\n\n".join(lines), 15, WHITE)
        _notes(s, "All sources were proposed by the Curator stage and "
                  "are checked by the Verifier before instructor review.")

    path = os.path.join(OUTPUT_DIR, f"{slug}-deck.pptx")
    prs.save(path)
    return path


def build_deck(session, plan, items):
    """Returns {file, slides} — raises RuntimeError via meta_api on API
    failure so the app's stage runner can log it."""
    spec = deck_spec(session, plan, items)
    topic = re.sub(r"[^a-z0-9]+", "-",
                   (session.get("topic") or "session").lower()).strip("-")
    slug = (topic[:40] or "session")
    path = render_pptx(session, spec, items, slug)
    return {"file": os.path.basename(path),
            "slides": len(spec.get("slides", [])) + 2}


CARD_SYSTEM = (
    "You are the Reference Card Builder stage of ClassPrep Autopilot. "
    "Build a ONE-PAGE in-class reference card for students: the distilled "
    "essentials of the session they can keep open during discussion and "
    "the lab. Dense but scannable. Return JSON only: {\"headline\": str, "
    "\"key_terms\": [{\"term\": str, \"definition\": str}], "
    "\"core_ideas\": [str], \"compare\": {\"left_label\": str, "
    "\"left\": [str], \"right_label\": str, \"right\": [str]}, "
    "\"try_it\": [str], \"discussion_prompt\": str}. 4-6 key terms, 3-5 "
    "core ideas, a two-column comparison when the topic has a natural "
    "contrast (else null), 2-3 try-it steps, one discussion prompt. "
    "No prose outside the JSON."
)


def build_card(session, plan, items):
    """Spark writes the card content; we render a printable one-page
    HTML card (light theme — students print these)."""
    raw = meta_api.chat(
        CARD_SYSTEM,
        "Session input:\n" + json.dumps(session, indent=2)
        + "\n\nPlan:\n" + (plan if isinstance(plan, str)
                           else json.dumps(plan, indent=2))
        + "\n\nCurated items:\n" + (items if isinstance(items, str)
                                    else json.dumps(items, indent=2)),
        max_output_tokens=4000)
    card = _as_json(raw)
    if not card:
        raise RuntimeError("Reference Card Builder: model returned no "
                           "parseable card JSON")
    topic = re.sub(r"[^a-z0-9]+", "-",
                   (session.get("topic") or "session").lower()).strip("-")
    slug = (topic[:40] or "session")
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    path = os.path.join(OUTPUT_DIR, f"{slug}-card.html")
    with open(path, "w") as f:
        f.write(_render_card_html(session, card, items))
    return {"file": os.path.basename(path),
            "key_terms": len(card.get("key_terms", []))}


def _esc(s):
    return (str(s or "").replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;"))


def _render_card_html(session, card, items):
    terms = "".join(
        f"<div class='term'><b>{_esc(t.get('term'))}</b> — "
        f"{_esc(t.get('definition'))}</div>"
        for t in card.get("key_terms", []))
    ideas = "".join(f"<li>{_esc(i)}</li>" for i in card.get("core_ideas", []))
    cmp_html = ""
    cmp_ = card.get("compare")
    if cmp_:
        left = "".join(f"<li>{_esc(x)}</li>" for x in cmp_.get("left", []))
        right = "".join(f"<li>{_esc(x)}</li>" for x in cmp_.get("right", []))
        cmp_html = (f"<h2>{_esc(cmp_.get('left_label'))} vs "
                    f"{_esc(cmp_.get('right_label'))}</h2>"
                    f"<div class='cols'><div><ul>{left}</ul></div>"
                    f"<div><ul>{right}</ul></div></div>")
    tryit = "".join(f"<li>{_esc(x)}</li>" for x in card.get("try_it", []))
    items = _items_list(items)
    sources = "".join(
        f"<div class='src'>{_esc(it.get('title'))} — "
        f"{_esc(it.get('source_name'))}: {_esc(it.get('source_url'))}</div>"
        for it in items)
    return f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<title>Reference Card — {_esc(session.get('topic'))}</title>
<style>
@page {{ size: letter; margin: 0.5in; }}
body {{ font-family: Georgia, 'Times New Roman', serif; color: #14161c;
  max-width: 7.5in; margin: 0 auto; font-size: 11.5pt; line-height: 1.35; }}
h1 {{ font-size: 17pt; margin: 0 0 2px; }}
.sub {{ color: #555; font-size: 10pt; margin-bottom: 8px; }}
h2 {{ font-size: 12pt; margin: 10px 0 4px; border-bottom: 1.5px solid #14161c;
  padding-bottom: 2px; }}
.term {{ margin: 3px 0; }}
ul {{ margin: 4px 0 4px 16px; padding: 0; }}
.cols {{ display: flex; gap: 18px; }} .cols > div {{ flex: 1; }}
.prompt {{ background: #eef2f9; border-left: 4px solid #3b6fd4;
  padding: 6px 10px; margin-top: 10px; }}
.src {{ font-size: 8.5pt; color: #444; margin: 2px 0; word-break: break-all; }}
.draft {{ color: #a33; font-size: 9pt; letter-spacing: 1px; }}
</style></head><body>
<div class="draft">DRAFT — BUILT BY CLASSPREP AUTOPILOT, AWAITING INSTRUCTOR REVIEW</div>
<h1>{_esc(card.get('headline') or session.get('topic'))}</h1>
<div class="sub">{_esc(session.get('course'))} · {_esc(session.get('date'))}</div>
<h2>Key terms</h2>{terms}
<h2>Core ideas</h2><ul>{ideas}</ul>
{cmp_html}
<h2>Try it</h2><ol>{tryit}</ol>
<div class="prompt"><b>Discuss:</b> {_esc(card.get('discussion_prompt'))}</div>
<h2>Sources</h2>{sources}
</body></html>"""


PODCAST_SYSTEM = (
    "You are the Podcast Builder stage of ClassPrep Autopilot. Write a "
    "prep podcast script for the INSTRUCTOR — a ~5 minute listen "
    "(750-850 words) they play while getting ready to teach this "
    "session. Single host, warm and brisk, spoken style. Structure with "
    "markdown section headers: ## Cold open (one hook sentence), "
    "## The session in one minute, ## The ideas that matter (walk the "
    "key concepts), ## Where students will push back (use the plan's "
    "likely student questions, with a one-line steer for each), "
    "## Running the room (timing cues from the session arc), "
    "## One-line close. Plain script text under each header — this is "
    "read aloud, so no bullet lists, no links, no stage directions "
    "beyond the headers. Return markdown only."
)


def build_podcast(session, plan, items):
    """Returns {file, words, est_minutes} for the script file."""
    script = meta_api.chat(
        PODCAST_SYSTEM,
        "Session input:\n" + json.dumps(session, indent=2)
        + "\n\nPlan:\n" + (plan if isinstance(plan, str)
                           else json.dumps(plan, indent=2))
        + "\n\nCurated items:\n" + (items if isinstance(items, str)
                                    else json.dumps(items, indent=2)),
        max_output_tokens=8000)  # long-form: Spark reasoning eats budget
    if not script or len(script) < 200:
        raise RuntimeError("Podcast Builder: model returned no usable "
                           "script")
    script = re.sub(r"(?m)^(## [^(\n]+)\s*\([^)]*\)\s*$", r"\1",
                    script)  # drop instruction parentheticals in headers
    words = len(re.findall(r"\S+", script))
    est = round(words / 150.0, 1)
    topic = re.sub(r"[^a-z0-9]+", "-",
                   (session.get("topic") or "session").lower()).strip("-")
    slug = (topic[:40] or "session")
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    path = os.path.join(OUTPUT_DIR, f"{slug}-podcast.md")
    header = (f"# Prep podcast — {session.get('topic') or 'Session'}\n\n"
              f"*{session.get('course') or ''} · "
              f"{session.get('date') or ''} · ~{est} min at speaking "
              f"pace ({words} words)*\n\n"
              f"DRAFT — built by ClassPrep Autopilot, awaiting "
              f"instructor review.\n\n---\n\n")
    with open(path, "w") as f:
        f.write(header + script.strip() + "\n")
    return {"file": os.path.basename(path), "words": words,
            "est_minutes": est}


DECK_SOURCE_SYSTEM = (
    "You are preparing a DECK SOURCE document for NotebookLM's Studio "
    "slide generator. One block per slide: a headline line, then 2-3 "
    "short bullet lines of on-slide content, then a line starting "
    "'Visual:' describing the image or layout. CRITICAL RULE: no "
    "narration, no speaker script, no 'say this' text anywhere — the "
    "slides must stand on their own as visuals that support a talk, "
    "not carry it. 7-9 slides following the session arc. Plain "
    "markdown, no preamble."
)


def _slug(session):
    topic = re.sub(r"[^a-z0-9]+", "-",
                   (session.get("topic") or "session").lower()).strip("-")
    return topic[:40] or "session"


def build_notebooklm_pack(session, plan, items):
    """Studio Agent (guided handoff): assembles a NotebookLM Source
    Pack as a zip — 01 sources, 02 deck source (narration-free, per
    the instructor's rule), 03 seeded chat prompts, 04 Studio recipe
    cards. NotebookLM has no consumer API, so the pack is built for
    one-click upload by the instructor; returned artifacts re-enter
    through the Verifier and Review gate like native outputs."""
    import zipfile

    plan_j = _as_json(plan) or {}
    items_j = _items_list(items)

    deck_source = meta_api.chat(
        DECK_SOURCE_SYSTEM,
        "Session input:\n" + json.dumps(session, indent=2)
        + "\n\nPlan:\n" + (plan if isinstance(plan, str)
                           else json.dumps(plan, indent=2)),
        max_output_tokens=6000)
    if not deck_source or len(deck_source) < 100:
        raise RuntimeError("NotebookLM pack: deck source generation "
                           "failed")

    sources_md = [f"# Sources — {session.get('topic') or 'Session'}",
                  f"\n{session.get('course') or ''} · "
                  f"{session.get('date') or ''}",
                  f"\n## Assigned for this session\n\n"
                  f"{session.get('assigned') or ''}",
                  "\n## Curated supporting items (proposed by the "
                  "Curator stage; checked by the Verifier)\n"]
    for it in items_j:
        sources_md.append(
            f"- **{it.get('title')}** ({it.get('type')}) — "
            f"{it.get('source_name')}: {it.get('source_url')}\n"
            f"  Why it fits: {it.get('why_it_fits')}")

    prompts = ["# Chat prompts — paste into NotebookLM Chat\n",
               "Source-grounded: answers cite only the uploaded "
               "sources.\n"]
    for q in (plan_j.get("likely_student_questions") or [])[:5]:
        prompts.append(f"- A student will ask: \"{q}\" — answer from "
                       f"the sources and note which source supports "
                       f"each claim.")
    for obj in (plan_j.get("objectives") or [])[:3]:
        prompts.append(f"- Explain, from the sources only: {obj}")

    recipes = """# Studio recipe cards

Upload `01-sources.md` and `02-deck-source.md` as sources in a new
NotebookLM notebook, then run any of these from the Studio panel.
NotebookLM output language is chosen in the Studio panel per artifact
— set it before generating (e.g. German for a language class, English
for a content class).

## Slide deck
Source: `02-deck-source.md` only. Generate the deck; each headline
becomes a slide. Review against the sources before class.

## Audio overview (~podcast)
Sources: `01-sources.md`. Choose the output language in the panel.
Good as a student-facing recap after the session.

## Flashcards / Quiz
Sources: `01-sources.md`. NotebookLM generates study sets grounded in
the uploaded items only — spot-check terms against the reference card.

## Mind map
Sources: `01-sources.md`. Use it to open the session: project the map
and ask students what is missing.

## Report / briefing doc
Sources: both files. A one-page briefing for a colleague covering the
same session.

Returned artifacts re-enter ClassPrep through the Verifier + Review
gate before use, like any native output.
"""

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    path = os.path.join(OUTPUT_DIR, f"{_slug(session)}-notebooklm.zip")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("01-sources.md", "\n".join(sources_md))
        z.writestr("02-deck-source.md",
                   f"# Deck source — narration-free (slides carry "
                   f"visuals only)\n\n{deck_source.strip()}\n")
        z.writestr("03-chat-prompts.md", "\n".join(prompts))
        z.writestr("04-studio-recipes.md", recipes)
    return {"file": os.path.basename(path), "files": 4}
